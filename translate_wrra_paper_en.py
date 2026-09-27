from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor


ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "output" / "WRRA_공간뒤틀림_은하원반면_암흑질량_우리은하모형_최원식_v1.0_2026-09-27.docx"
OUTDIR = ROOT / "output" / "english"
OUTDIR.mkdir(parents=True, exist_ok=True)
TARGET = OUTDIR / "WRRA_Spatial_Twist_Galactic_Disk_Dark_Mass_Milky_Way_Model_Wonsik_Choi_v1.0_EN_2026-09-28.docx"

FONT = "Liberation Sans"
INK = "000000"
MID = "5B6573"


PARAGRAPHS = {
    0: "WRRA Spatial Twist Selection of Galactic Disk Planes\nand the Dark Mass Phenotype",
    1: "A Dual Output Model of the Milky Way Warp and Rotation Curve",
    2: "Wonsik Choi\nIndependent Researcher  Seoul Republic of Korea\njanefather@gmail.com\n28 September 2026  English Version v1.0",
    3: "Research status  Quantitative toy model and falsifiable WRRA hypothesis",
    5: "Abstract",
    6: "This paper formulates the selection of an approximately two-dimensional galactic disk plane in three-dimensional space as a problem of WRRA spatial twist. The central revision is to separate the roles of twist and angular momentum. Twist neither creates rotation nor determines its direction. The directional sector of twist selects a local plane field that does not distinguish up from down, while its magnitude sector is stored as spatial stress and generates additional gravity that appears as dark mass. Angular momentum is a separate physical quantity that subsequently maintains rotation within the selected plane and prevents central collapse.",
    7: "This distinction is applied to the Milky Way. A thin-disk approximation to the McMillan baryonic mass model is combined with the fixed WRRA transition scale aₜ=1.1918×10⁻¹⁰ m s⁻² and a single constitutive response calibrated with SPARC. Without galaxy-specific recalibration, the resulting rotation speed differs from a linearized version of the Eilers et al. 5-25 kpc rotation curve by an RMS of 2.42 km s⁻¹ over 8-25 kpc. The model gives 226.3 km s⁻¹ at the solar radius of 8.2 kpc and 201.6 km s⁻¹ at 25 kpc. The disk plane is calibrated to the onset of the Gaia DR3 Cepheid warp at R≈11 kpc, an outer inclination of about 3 degrees, and the radial change of the line of nodes. These results do not prove the twist hypothesis. They provide the first closed model in which one twist state produces disk shape and additional gravity as distinct outputs.",
    8: "Keywords  WRRA  spatial twist  galactic disk  Milky Way  warp  dark matter mass phenotype  rotation curve",
    9: "Core Assessment",
    11: "1 Research Question",
    12: "The existence of angular momentum in a galactic disk does not by itself answer why matter gathered into that plane. Standard gas cooling and angular-momentum conservation are powerful mechanisms for flattening and maintaining a disk. WRRA asks a prior geometrical question: how is a particular two-dimensional transport plane selected in three-dimensional space, and why does that plane bend and twist with position?",
    13: "This paper proposes that spatial twist first selects an unoriented plane, after which matter and angular momentum align with and operate within that plane. Here unoriented means that up and down are not assigned. A plane has two opposite normals, but both represent the same plane. A local disk plane can therefore be defined without introducing a privileged global vertical direction.",
    14: "The second question is whether another sector of the same twist state participates in dark matter phenomena. The existing WRRA twist-stress model separates the direct gravity of ordinary matter gₘ from the additional gravity of spatial stress gₜ [1,2]. This study places the scalar magnitude sector and the directional sector that selects the disk plane within one twist state, without forcing them to be identical or governed by the same equation.",
    15: "2 Elements Preserved from the Existing WRRA Model",
    16: "This paper does not alter WRRA Core or Minimal Computation Cosmology 2.3.2. Standard physical descriptions of ordinary matter and the four interactions are retained. Instead of immediately declaring a fifth force, additional gravity is calculated in the weak-field Renderer as an effective source in which twist stress stored in space is expressed as mass. The source branch of general relativity and the modified-response branch are not counted twice in the same calculation.",
    17: "Earlier work normalized the present-universe twist stress with H₀=67.4 km s⁻¹ Mpc⁻¹ and fₜ=0.265, giving aₜ=1.1918×10⁻¹⁰ m s⁻². Across the 201 lowest-acceleration SPARC points, this value differed from the diagnostic inverse estimate by about 0.75 percent. The present paper transfers that normalization to the Milky Way without refitting it.",
    18: "The new element is a computational layer for disk-plane selection. A scalar χₜ or ν that represents only the magnitude of twist stress cannot select a plane. The minimal state must therefore record a plane director that has direction but no sign, together with energy density and anisotropic stress.",
    19: "3 Separating the Roles of Twist and Angular Momentum",
    20: "3.1 Unoriented Plane",
    21: "Let n be a unit normal representing a local plane. Because n and -n represent the same plane, the physical state is the equivalence class [n], not the vector itself. The plane projector P is invariant under this sign reversal.",
    23: "This structure does not specify up and down in advance. It also does not specify clockwise or counterclockwise rotation. The twist plane field therefore neither defines angular momentum nor hides the sign of angular momentum within its definition.",
    24: "3.2 Shape of the Surface",
    25: "The mean disk plane is represented as a level set S=0. With this definition, position-dependent normals construct one continuous surface.",
    27: "For a general normal field to integrate into a physical surface, the Frobenius integrability condition must hold. Because this paper constructs the normal from S, the condition is satisfied by construction.",
    29: "3.3 Magnitude Sector of the Same Twist State",
    30: "To separate the directional sector of twist from its gravitational magnitude sector, the effective stress contains an energy density ρₜ, isotropic pressure pₜ, and anisotropic stress πₜ defined relative to the unoriented normal. The sign and magnitude of qₜ represent the difference between in-plane and out-of-plane responses.",
    32: "The quantity ρₜ controls the magnitude of additional gravity, whereas [n] and qₜ control the orientation of the disk plane and its vertical restoring response. The sectors belong to one state but are not the same number. This is the minimal decomposition that turns the statement twist selects the plane but does not define angular momentum into a calculable model.",
    33: "4 Constitutive Equation for the Gravitational Sector",
    34: "4.1 Matter Gravity and Twist Stress Gravity",
    35: "For weak-field circular orbits, the observed gravitational acceleration is decomposed directly into a baryonic component and a spatial-stress component.",
    37: "An empirical response that is consistent with the full SPARC range is adopted and frozen as the calibrated constitutive equation of this model. This is a model revision that promotes a relation previously used as external support into the actual Milky Way calculation. No claim is made that the functional form has been derived from a WRRA action.",
    39: "The transition scale is fixed by the cosmological normalization.",
    41: "4.2 Source Formulation and Prohibition of Double Counting",
    42: "This paper uses the weak-field branch in which ρₜ,eff is read as an additional effective source. The same effect is not added again as a modified-geometry term. A covariant completion would require conservation of the total stress-energy tensor.",
    44: "The quantity ρₜ,eff is not a particle number density. It is the mass phenotype obtained when the same additional gravitational field is inverted as a Newtonian mass source. The equivalent twist mass in this paper must therefore not be read as a measurement of actual particle mass.",
    45: "5 Milky Way Validation Inputs",
    46: "The baryonic inputs are taken from the McMillan mass model of the Milky Way [3]. They include the central surface densities and scale lengths of the thin and thick stellar disks, the masses and scale lengths of HI and H₂ gas, and the bulge mass. The original model includes finite thickness and central gas holes, whereas the present rotation calculation approximates each component as a thin exponential disk. This simplification cannot describe the bar and bulge region inside 5 kpc precisely.",
    48: "The rotation-curve comparison uses the 5-25 kpc result from the analysis of more than 23,000 red giant stars by Eilers et al. [4]. That study reported 229.0 km s⁻¹ at the solar radius and an outer slope of -1.7 km s⁻¹ kpc⁻¹. The present paper uses a linearized target built from those two reported values rather than the full individual data set. The RMS values below are therefore first-order shape comparisons, not observational likelihood tests.",
    49: "Warp inputs are taken from Gaia DR3 Cepheid studies. Dehnen et al. reported no clear warp inside about 11 kpc, an inclination that increases outward to about 3 degrees, and a line of nodes that changes with radius [5]. Cabrera-Gadea et al. reported m=1 and m=2 components, a twisted line of nodes for R>11 kpc, and an m=1 pattern speed of about 9.2±3.1 km s⁻¹ kpc⁻¹ [6]. Low-order m=0,1,2 modes also dominate the HI disk [7].",
    50: "6 Calculation Method",
    51: "6.1 Baryonic Rotation Field",
    52: "The circular speed of each exponential disk is calculated with the standard thin-disk expression containing modified Bessel functions.",
    54: "The bulge is represented by a Hernquist approximation with mass Mᵇ and scale aᵇ.",
    56: "The squared velocities of the components are added to obtain gₘ=vbar²/R. The total circular speed is then calculated using the fixed ν and aₜ.",
    58: "6.2 Twist Plane and Warp",
    59: "The observed mean disk height can be expressed in low-order Fourier modes. This is the standard form used in existing warp analyses [5-7]. WRRA interprets it not merely as the posterior shape of a material disk but as the observational representation of the twist plane field S=0.",
    61: "The quantitative demonstration retains only the fundamental m=1 mode. The warp begins at 11 kpc, reaches 3 degrees at 14 kpc, and then saturates. The line of nodes changes by about 14 degrees per kpc over the observed interval and is conservatively held fixed beyond 16.5 kpc.",
    63: "Excluding the m=0 and m=2 modes does not deny their existence. The purpose is first to show how much can be reconstructed with one directional field, while leaving the lopsided warp as a structured residual for the next validation stage.",
    64: "7 Outputs",
    65: "7.1 Rotation Curve",
    67: "Over 8-25 kpc, the RMS difference between the WRRA curve and the linearized Eilers target is 2.42 km s⁻¹, the mean difference is -2.14 km s⁻¹, and the maximum absolute difference is 3.31 km s⁻¹. Over the full 5-25 kpc range, the RMS is 3.64 km s⁻¹. The largest discrepancy occurs near the inner boundary, where the bar and non-axisymmetry are strongest.",
    69: "Figure 1  Milky Way rotation curve calculated from the McMillan baryonic inputs and the fixed WRRA constitutive response  The dashed line is the linearized observational target from Eilers et al.",
    70: "Neither aₜ nor the functional form of ν was readjusted to fit the Milky Way rotation curve. The baryonic model itself is based on Milky Way observations, however, and ν is a SPARC-calibrated constitutive relation, so this is not a model-independent blind prediction. Under the WRRA assessment rule, it is classified as a Milky Way output obtained by freezing validated inputs and applying the same transformation.",
    71: "7.2 Equivalent Twist Mass",
    72: "The additional acceleration can be converted into a spherically equivalent Newtonian mass as follows.",
    74: "This equivalent mass is 3.41×10¹⁰ M☉ at 8.2 kpc, 1.25×10¹¹ M☉ at 20 kpc, and 1.63×10¹¹ M☉ at 25 kpc. If the physical twist stress has a two-dimensional or anisotropic distribution, its energy-density profile is not identical to this spherical equivalent mass. The quantity is a ledger of the additional gravity required by the rotation field.",
    75: "7.3 Disk Plane and Warp Shape",
    78: "Figure 2  Minimal m=1 realization of the WRRA twist plane using the Gaia Cepheid warp as a calibration input  The dotted circle marks the 11 kpc warp onset radius",
    79: "This plane was not constructed by entering an angular-momentum vector. The observed warp shape is used to fix the twist plane field, after which the degree to which stellar and gas orbits align with that plane remains a separate dynamical problem. The current figure shows only the mean plane and does not include disk thickness, m=2 asymmetry, spiral arms, or the bar.",
    80: "7.4 Two Outputs from One Twist State",
    81: "The minimal twist state in this paper is the tuple consisting of an unoriented plane [n], stress energy density ρₜ, and anisotropic magnitude qₜ. It produces the disk plane, additional gravity, and vertical restoring force as distinct Renderer outputs.",
    83: "Testing vertical motion in the disk requires a restoring potential for the distance S away from the plane. At lowest order it may be written as follows.",
    85: "The model is complete only when νᴢ can be calculated from ρₜ and qₜ fixed by the rotation curve. At present, no constitutive equation for qₜ has been specified, so hᴢ is not predicted numerically. It is important not to hide this gap.",
    86: "8 Distinctive WRRA Explanation and the Role of Existing Physics",
    88: "WRRA does not remove angular-momentum conservation, gas cooling, or the collisionless motion of stars. Its originality does not lie in inventing already known numerical values. It lies in separating and connecting plane selection and the dark-mass phenotype as directional and magnitude sectors of one twist state. Even when another theory uses the same rotation curve or warp equation, outputs calculated through the WRRA state and transformation remain WRRA outputs.",
    89: "Conversely, merely citing an external warp equation or RAR constitutive relation without calculating it inside WRRA would not count as integration. This paper includes both relations as calibrated model components and calculates Milky Way outputs, so they are integrated into the numerical model. Deriving those constitutive relations from a deeper WRRA action remains an open task.",
    90: "9 Falsification Conditions and Observation Plan",
    91: "The hypothesis is not validated merely because a disk is observed. Standard cooling and tidal torques can also produce disks and warps. The discriminatory power of WRRA comes from joint constraints left by the same twist state across different observables.",
    92: "If gₜ and ρₜ,eff fixed by the rotation curve cannot reproduce the stellar vertical velocity dispersion and disk thickness, the directional-stress coupling fails.",
    93: "If the twist stress required by dynamics cannot reproduce weak and strong gravitational lensing through the same metric, the dark-mass phenotype interpretation fails.",
    94: "If the warp plane zₜ cannot connect the mean planes and vertical velocity fields of stars, HI, and H₂ through one time evolution, the plane-field interpretation fails.",
    95: "If the m=2 lopsidedness and line-of-nodes variation are fully explained by instantaneous torques from external satellites, with no residual left for a fixed spatial-stress directional field, the upper bound on the additional WRRA term approaches zero.",
    96: "If polar-ring galaxies or mutually perpendicular multiple disks cannot be represented by one local plane field and arbitrary extra planes must be introduced, the minimal-computation hypothesis is weakened.",
    97: "If aₜ or ν must be freely refitted for each galaxy, the universal constitutive equation for spatial stress is falsified.",
    98: "The most direct next calculation is to solve the vertical Jeans equation within the same Milky Way baryonic model and constrain qₜ. After fixing gₜ from the rotation curve, the model must simultaneously fit the solar-neighborhood K_z, outer-disk flaring, Cepheid V_z, and HI thickness. This test determines whether the disk plane and dark gravity are truly two outputs of the same twist state.",
    99: "10 Limitations",
    100: "The m=1 warp shape is a calibrated input fitted to Gaia observations. It was not derived from first principles.",
    101: "The empirical SPARC relation ν is integrated into and calculated within this model, but it has not been uniquely derived from a twist action.",
    102: "The baryonic model simplifies the McMillan model into thin axisymmetric disks. It omits central gas holes, the bar, spiral arms, and non-circular motion.",
    103: "The Eilers target is a curve linearized from two reported values. It is not a likelihood test using the individual data points and their covariance.",
    104: "The equivalent twist mass is a spherical inverse estimate and is not the physical mass distribution of a two-dimensional stress energy field.",
    105: "The constitutive equation for qₜ and the vertical restoring frequency νᴢ have not yet been fixed, so the model does not predict disk thickness.",
    106: "A covariant completion that simultaneously calculates structure growth, the CMB, cluster collisions, and lensing lies outside the present scope.",
    107: "11 Conclusion",
    108: "This paper presents a calculable WRRA hypothesis that differs from the explanation that angular momentum creates the disk plane. The directional sector of spatial twist first selects a plane that does not distinguish up from down, while angular momentum maintains rotation within that plane. An initially global vertical direction or a two-dimensional rotation plane therefore need not be supplied in order for an approximately two-dimensional disk to arise in three-dimensional space.",
    109: "The magnitude sector of twist produces additional gravity gₜ under the existing WRRA model and appears as the dark-mass phenotype. Using the Milky Way baryonic model with fixed aₜ and fixed ν, the calculated curve differs from the linearized observed rotation curve by an RMS of 2.42 km s⁻¹ over 8-25 kpc. The directional sector implements the observed warp beyond 11 kpc and an inclination of about 3 degrees as one continuous plane.",
    110: "The present result does not identify the disk plane and dark gravity with one equation. Instead, it separates and calculates them as two outputs from one twist state. The final assessment is that the first closure of the quantitative Milky Way toy model succeeds, while the claim that spatial twist is the physical cause of the disk plane remains unconfirmed. The next falsification stage is to determine whether the stress field fixed by the rotation curve also passes vertical-equilibrium and lensing tests without additional calibration.",
    111: "Appendix A Reproducible Values",
    113: "Appendix B Claim Grades",
    115: "References",
}


TABLES = [
    [
        ["Stage", "Content", "Current assessment"],
        ["Validation inputs", "Milky Way baryonic model and rotation curve  Gaia warp  H₀  fₜ", "External observations"],
        ["WRRA transformation", "Unoriented plane field [n] and projector P  constitutive twist-stress response ν", "Calculated in this paper"],
        ["Outputs", "Disk plane zₜ  gₜ  vWRRA  equivalent twist mass", "Quantitative outputs"],
        ["Falsification condition", "The same twist state fails to pass warp  rotation  vertical equilibrium and lensing jointly", "Untested"],
        ["Conclusion", "The first Milky Way toy model works numerically but the causal hypothesis remains unconfirmed", "Conditional pass"],
    ],
    [
        ["Component", "Mass 10⁹ M☉", "Scale kpc", "Representation"],
        ["Thin stellar disk", "35.66", "2.53", "Exponential disk"],
        ["Thick stellar disk", "11.25", "3.38", "Exponential disk"],
        ["HI gas", "11.00", "7.00", "Exponential disk"],
        ["H₂ gas", "1.20", "1.50", "Exponential disk"],
        ["Bulge", "9.13", "a=0.70", "Hernquist approximation"],
    ],
    [
        ["R kpc", "Baryonic v", "gₜ/gₘ", "WRRA v", "Target v", "Mₜ,eff 10¹⁰ M☉"],
        ["5.0", "192.4", "0.319", "221.0", "234.4", "1.37"],
        ["8.2", "182.6", "0.537", "226.3", "229.0", "3.41"],
        ["12.0", "161.0", "0.868", "220.1", "222.5", "6.28"],
        ["16.5", "138.9", "1.320", "211.6", "214.9", "9.77"],
        ["20.0", "125.9", "1.692", "206.6", "208.9", "12.48"],
        ["25.0", "112.0", "2.237", "201.6", "200.4", "16.32"],
    ],
    [
        ["R kpc", "Inclination deg", "Line of nodes deg", "m=1 maximum height kpc"],
        ["5.0", "0.00", "Undefined", "0.000"],
        ["8.2", "0.00", "Undefined", "0.000"],
        ["12.0", "1.00", "-1.0", "0.209"],
        ["16.5", "3.00", "62.0", "0.865"],
        ["20.0", "3.00", "62.0", "1.048"],
        ["25.0", "3.00", "62.0", "1.310"],
    ],
    [
        ["Question", "Standard physical explanation", "Connection added by WRRA"],
        ["Why is the disk flat", "Gas collisions and radiative cooling reduce vertical motion", "The twist plane supplies the reference plane relative to which vertical residual motion is removed"],
        ["Why that plane", "A plane perpendicular to the total angular-momentum vector", "An unoriented transport plane [n] is hypothesized to precede angular momentum"],
        ["Why does the outer disk warp", "Satellite interactions and halo torques", "A position-dependent plane field zₜ and anisotropic stress produce the warp"],
        ["Why is there additional gravity", "A particle dark matter halo or modified gravity", "The mass phenotype gₜ and ρₜ,eff of spatial stress"],
    ],
    [
        ["Item", "Value", "Status"],
        ["H₀", "67.4 km s⁻¹ Mpc⁻¹", "Validation input"],
        ["fₜ", "0.265", "Validation input"],
        ["aₜ", "1.191812669×10⁻¹⁰ m s⁻²", "WRRA closure output"],
        ["Stellar disk mass", "46.910×10⁹ M☉", "Converted McMillan input"],
        ["Gas mass", "12.2×10⁹ M☉", "McMillan input"],
        ["Bulge mass", "9.13×10⁹ M☉", "McMillan input"],
        ["WRRA vc at 8.2 kpc", "226.34 km s⁻¹", "Output"],
        ["WRRA vc at 25 kpc", "201.59 km s⁻¹", "Output"],
        ["Mₜ,eff at 25 kpc", "1.632×10¹¹ M☉", "Spherical equivalent output"],
        ["Warp onset", "11 kpc", "Observational calibration input"],
        ["Outer inclination", "3°", "Observational calibration input"],
    ],
    [
        ["Claim", "Grade", "Basis for assessment"],
        ["An unoriented plane is defined without up or down", "Mathematical fact", "P is invariant under n→-n"],
        ["The twist plane precedes the material disk plane", "WRRA hypothesis", "Requires cross-validation with vertical equilibrium"],
        ["Twist stress generates additional gravity", "WRRA causal hypothesis", "First rotation test passed  lensing and growth untested"],
        ["Reproduction of the Milky Way rotation curve", "Post-calibration output", "Calculated from fixed aₜ and ν with baryonic inputs"],
        ["Reconstruction of the Milky Way warp surface", "Observationally calibrated model", "Gaia inclination and line-of-nodes variation used as inputs"],
        ["Common cause of the two outputs", "Integration hypothesis", "qₜ and νᴢ remain unspecified"],
    ],
]


def set_run_font(run, size=None, bold=None, italic=None, color=INK):
    run.font.name = FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def clear_runs(paragraph):
    for run in list(paragraph.runs):
        paragraph._p.remove(run._r)


def replace_paragraph(paragraph, text, size=None, bold=None, italic=None, color=INK):
    clear_runs(paragraph)
    run = paragraph.add_run(text.replace("–", "-").replace("—", "-"))
    set_run_font(run, size=size, bold=bold, italic=italic, color=color)
    return run


def translate_document():
    doc = Document(SOURCE)

    for style_name in ["Normal", "Body Text", "Title", "Heading 1", "Heading 2", "Heading 3", "List Bullet"]:
        if style_name in doc.styles:
            style = doc.styles[style_name]
            style.font.name = FONT
            style._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT)
            style._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT)
            style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
            style.font.color.rgb = RGBColor.from_string(INK)
    doc.styles["Body Text"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

    for index, text in PARAGRAPHS.items():
        paragraph = doc.paragraphs[index]
        style_name = paragraph.style.name
        if index == 0:
            replace_paragraph(paragraph, text, size=22, bold=True)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif index == 1:
            replace_paragraph(paragraph, text, size=14)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif index == 2:
            clear_runs(paragraph)
            lines = text.split("\n")
            for line_index, line in enumerate(lines):
                run = paragraph.add_run(line)
                set_run_font(run, size=12 if line_index == 0 else 10.5, bold=(line_index == 0))
                if line_index < len(lines) - 1:
                    run.add_break()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif index == 3:
            replace_paragraph(paragraph, text, size=10.5, italic=True, color=MID)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif index in (69, 78):
            replace_paragraph(paragraph, text, size=9, italic=True, color=MID)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif style_name == "Heading 1":
            replace_paragraph(paragraph, text, size=16, bold=True)
        elif style_name == "Heading 2":
            replace_paragraph(paragraph, text, size=13, bold=True)
        elif style_name == "List Bullet":
            replace_paragraph(paragraph, text, size=10.2)
        else:
            replace_paragraph(paragraph, text, size=10.5)

    for index in range(116, 128):
        paragraph = doc.paragraphs[index]
        replace_paragraph(paragraph, paragraph.text, size=8.5)

    doc.paragraphs[113].paragraph_format.page_break_before = True

    for table_index, translated_rows in enumerate(TABLES):
        table = doc.tables[table_index]
        for row_index, translated_cells in enumerate(translated_rows):
            for col_index, text in enumerate(translated_cells):
                paragraph = table.rows[row_index].cells[col_index].paragraphs[0]
                clear_runs(paragraph)
                run = paragraph.add_run(text.replace("–", "-").replace("—", "-"))
                set_run_font(run, size=8.4 if table_index in (0, 4, 6) else 8.7, bold=(row_index == 0), color="FFFFFF" if row_index == 0 else INK)

    core = doc.core_properties
    core.title = "WRRA Spatial Twist Selection of Galactic Disk Planes and the Dark Mass Phenotype"
    core.subject = "A dual output model of the Milky Way warp and rotation curve"
    core.author = "Wonsik Choi"
    core.keywords = "WRRA, spatial twist, galactic disk, Milky Way, warp, dark matter phenotype"
    doc.save(TARGET)
    return TARGET


if __name__ == "__main__":
    print(translate_document())
