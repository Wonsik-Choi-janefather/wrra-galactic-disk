from __future__ import annotations

import math
import os
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
from matplotlib import cm
from scipy.special import iv, kv

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
OUT = ROOT / "output"
FIG = OUT / "figures"
EQ = OUT / "equations"
OUT.mkdir(exist_ok=True)
FIG.mkdir(exist_ok=True)
EQ.mkdir(exist_ok=True)

DOCX_PATH = OUT / "WRRA_공간뒤틀림_은하원반면_암흑질량_우리은하모형_최원식_v1.0_2026-09-27.docx"

FONT = "Noto Sans CJK KR"
INK = "000000"
NAVY = "1F4E78"
PALE = "EAF2F8"
LIGHT = "F4F7FA"
GRID = "D9D9D9"
MID = "5B6573"

G = 4.30091e-6  # kpc (km/s)^2 / solar mass
ACC_CONV = 1e6 / 3.0856775814913673e19  # [(km/s)^2/kpc] -> m/s^2
A_T = 1.191812669e-10


INLINE_SYMBOLS = {
    "M_T^eff": "Mₜ,eff",
    "ρ_T^eff": "ρₜ,eff",
    "v_WRRA": "vWRRA",
    "v_bary": "vbar",
    "v_c": "vc",
    "g_obs": "gobs",
    "g_M": "gₘ",
    "g_T": "gₜ",
    "a_T": "aₜ",
    "f_T": "fₜ",
    "ρ_T": "ρₜ",
    "q_T": "qₜ",
    "χ_T": "χₜ",
    "z_T": "zₜ",
    "h_z": "hᴢ",
    "ν_z": "νᴢ",
    "σ_z": "σᴢ",
    "R_d": "Rᴅ",
    "M_b": "Mᵇ",
    "a_b": "aᵇ",
    "Z_0": "Z₀",
    "Z_1": "Z₁",
    "Z_2": "Z₂",
    "φ_1": "φ₁",
    "φ_2": "φ₂",
}


def clean_inline(text):
    text = str(text)
    for source, target in INLINE_SYMBOLS.items():
        text = text.replace(source, target)
    return text


def exp_disc_mass(sigma0_msun_pc2: float, rd_kpc: float) -> float:
    sigma0 = sigma0_msun_pc2 * 1e6
    return 2.0 * math.pi * sigma0 * rd_kpc**2


COMPONENTS = {
    "thin": {"mass": exp_disc_mass(886.7, 2.53), "rd": 2.53},
    "thick": {"mass": exp_disc_mass(156.7, 3.38), "rd": 3.38},
    "HI": {"mass": 1.1e10, "rd": 7.0},
    "H2": {"mass": 1.2e9, "rd": 1.5},
}
M_BULGE = 9.13e9
A_BULGE = 0.7


def v2_exp_disc(r, mass, rd):
    r = np.asarray(r, dtype=float)
    y = r / (2.0 * rd)
    kernel = iv(0, y) * kv(0, y) - iv(1, y) * kv(1, y)
    return 2.0 * G * mass / rd * y**2 * kernel


def model_values(r):
    r = np.asarray(r, dtype=float)
    v2_bary = sum(v2_exp_disc(r, c["mass"], c["rd"]) for c in COMPONENTS.values())
    v2_bary += G * M_BULGE * r / (r + A_BULGE) ** 2
    v_bary = np.sqrt(v2_bary)
    g_m = (v2_bary / r) * ACC_CONV
    y = g_m / A_T
    nu = 1.0 / (1.0 - np.exp(-np.sqrt(y)))
    g_total = nu * g_m
    g_t = g_total - g_m
    v_total = np.sqrt((g_total / ACC_CONV) * r)
    m_t_eff = (g_t / ACC_CONV) * r**2 / G
    return v_bary, g_m, nu, g_t, v_total, m_t_eff


def eilers_linear(r):
    return 229.0 - 1.7 * (np.asarray(r) - 8.2)


def warp_parameters(r):
    r = np.asarray(r, dtype=float)
    inc_deg = 3.0 * np.clip((r - 11.0) / 3.0, 0.0, 1.0)
    phi_deg = -15.0 + 14.0 * np.clip(r - 11.0, 0.0, 5.5)
    amp = r * np.tan(np.deg2rad(inc_deg))
    return inc_deg, phi_deg, amp


def warp_height(r, phi):
    inc, phi1, amp = warp_parameters(r)
    return amp * np.sin(phi - np.deg2rad(phi1))


def create_equations():
    equations = {
        "eq01_projector": r"\mathbf{n}\sim-\mathbf{n},\qquad P_{ij}=\delta_{ij}-n_i n_j,\qquad P^2=P",
        "eq02_surface": r"S(R,\phi,z,t)=z-z_T(R,\phi,t)=0,\qquad \mathbf{n}=\frac{\nabla S}{\|\nabla S\|}",
        "eq03_integrability": r"\mathbf{n}\cdot(\nabla\times\mathbf{n})=0",
        "eq04_stress": r"T_T^{\mu\nu}=\rho_T u^\mu u^\nu+p_T h^{\mu\nu}+\pi_T^{\mu\nu},\qquad \pi_{T,ij}=q_T\left(n_i n_j-\frac{1}{3}\delta_{ij}\right)",
        "eq05_gravity": r"g_{\mathrm{obs}}(R)=g_M(R)+g_T(R)=\nu\!\left(\frac{g_M}{a_T}\right)g_M(R)",
        "eq06_response": r"\nu(y)=\frac{1}{1-\exp(-\sqrt{y})},\qquad g_T=[\nu(y)-1]g_M",
        "eq07_scale": r"a_T=cH_0\sqrt{\frac{f_T}{8}}=1.1918\times10^{-10}\ \mathrm{m\,s^{-2}}",
        "eq08_disk": r"v_d^2(R)=4\pi G\Sigma_0R_d y^2\!\left[I_0(y)K_0(y)-I_1(y)K_1(y)\right],\qquad y=\frac{R}{2R_d}",
        "eq09_bulge": r"v_b^2(R)=\frac{GM_bR}{(R+a_b)^2}",
        "eq10_warp": r"z_T=Z_0(R)+Z_1(R)\sin[\phi-\phi_1(R,t)]+Z_2(R)\sin\!\left(2[\phi-\phi_2(R,t)]\right)",
        "eq11_baseline_warp": r"Z_1(R)=R\tan i(R),\quad i(R)=3^\circ\,\mathrm{clip}\!\left(\frac{R-11}{3},0,1\right),\quad \phi_1(R)=-15^\circ+14^\circ\,\mathrm{clip}(R-11,0,5.5)",
        "eq12_velocity": r"v_{\mathrm{WRRA}}(R)=\sqrt{R\,g_{\mathrm{obs}}(R)}",
        "eq13_mass": r"M_T^{\mathrm{eff}}(<R)=\frac{R^2g_T(R)}{G}",
        "eq14_balance": r"\nabla^2\Phi=4\pi G\left(\rho_b+\rho_T^{\mathrm{eff}}\right),\qquad \nabla_\mu\left(T_b^{\mu\nu}+T_T^{\mu\nu}\right)=0",
        "eq15_thickness": r"\Phi_\perp\simeq\frac{1}{2}\nu_z^2S^2,\qquad h_z\simeq\frac{\sigma_z}{\nu_z}",
        "eq16_chain": r"\mathcal{T}=([\mathbf{n}],\rho_T,q_T)\quad\Longrightarrow\quad\{z_T(R,\phi,t),g_T(R),\Phi_\perp\}",
    }
    for name, expr in equations.items():
        width = 8.2 if len(expr) > 115 else (7.2 if len(expr) > 75 else 5.2)
        fig = plt.figure(figsize=(width, 0.72), dpi=300, facecolor="white")
        fig.text(0.5, 0.5, f"${expr}$", ha="center", va="center", fontsize=15, color="black")
        fig.savefig(EQ / f"{name}.png", dpi=300, bbox_inches="tight", pad_inches=0.06, facecolor="white")
        plt.close(fig)


def create_figures():
    plt.rcParams.update({
        "font.family": "DejaVu Sans",
        "font.size": 9.5,
        "axes.titlesize": 11,
        "axes.labelsize": 9.5,
        "legend.fontsize": 8.5,
        "figure.dpi": 160,
    })

    r = np.linspace(5.0, 25.0, 401)
    v_b, g_m, nu, g_t, v_w, m_eff = model_values(r)
    target = eilers_linear(r)

    fig, axes = plt.subplots(2, 1, figsize=(7.2, 6.2), sharex=True,
                             gridspec_kw={"height_ratios": [2.2, 1.0], "hspace": 0.08})
    ax = axes[0]
    ax.plot(r, target, color="#1f1f1f", lw=2.0, ls="--", label="Eilers linearized target")
    ax.plot(r, v_b, color="#7f8c8d", lw=2.0, label="Baryons only")
    ax.plot(r, v_w, color="#1f77b4", lw=2.4, label="WRRA calibrated closure")
    ax.axvline(8.2, color="#aeb6bf", lw=1.0, ls=":")
    ax.set_ylabel("Circular speed  km s$^{-1}$")
    ax.set_xlim(5, 25)
    ax.set_ylim(95, 250)
    ax.grid(alpha=0.22)
    ax.legend(loc="lower left", frameon=False, ncol=3)
    ax.set_title("Milky Way rotation curve in the WRRA twist-stress closure")

    ax2 = axes[1]
    ax2.plot(r, nu - 1.0, color="#c0392b", lw=2.1, label="$g_T/g_M$")
    ax2.set_ylabel("Twist / direct")
    ax2.set_xlabel("Galactocentric radius  kpc")
    ax2.set_ylim(0, 2.55)
    ax2.grid(alpha=0.22)
    ax2b = ax2.twinx()
    ax2b.plot(r, m_eff / 1e11, color="#7d3c98", lw=1.8, label="$M_T^{eff}$")
    ax2b.set_ylabel(r"Spherical-equivalent $M_T$  $10^{11} M_\odot$")
    handles = ax2.get_lines() + ax2b.get_lines()
    labels = [h.get_label() for h in handles]
    ax2.legend(handles, labels, loc="upper left", frameon=False, ncol=2)
    fig.savefig(FIG / "rotation_curve.png", bbox_inches="tight", facecolor="white")
    plt.close(fig)

    rr = np.linspace(4.0, 25.0, 180)
    pp = np.linspace(0, 2 * np.pi, 240)
    R, P = np.meshgrid(rr, pp)
    X = R * np.cos(P)
    Y = R * np.sin(P)
    Z = warp_height(R, P)

    fig = plt.figure(figsize=(7.4, 6.0))
    ax3 = fig.add_subplot(121, projection="3d")
    norm = plt.Normalize(-1.4, 1.4)
    ax3.plot_surface(X, Y, Z, rstride=4, cstride=5,
                     facecolors=cm.coolwarm(norm(Z)), linewidth=0, antialiased=True, shade=False)
    ax3.set_xlim(-25, 25)
    ax3.set_ylim(-25, 25)
    ax3.set_zlim(-2.2, 2.2)
    ax3.set_xlabel("x  kpc", labelpad=2)
    ax3.set_ylabel("y  kpc", labelpad=2)
    ax3.set_zlabel("z  kpc", labelpad=2)
    ax3.set_title("Twist-selected surface")
    ax3.view_init(elev=24, azim=-56)
    ax3.set_box_aspect((1, 1, 0.24))

    ax4 = fig.add_subplot(122)
    pcm = ax4.pcolormesh(X, Y, Z, shading="auto", cmap="coolwarm", vmin=-1.4, vmax=1.4)
    ax4.add_patch(plt.Circle((0, 0), 11, fill=False, color="black", ls=":", lw=1.0))
    ax4.set_aspect("equal")
    ax4.set_xlim(-25, 25)
    ax4.set_ylim(-25, 25)
    ax4.set_xlabel("x  kpc")
    ax4.set_ylabel("y  kpc")
    ax4.set_title("Face-on height map")
    cbar = fig.colorbar(pcm, ax=ax4, shrink=0.72, pad=0.04)
    cbar.set_label("z  kpc")
    fig.suptitle("Minimal m=1 Milky Way warp used as the twist-plane calibration", y=0.99, fontsize=11)
    fig.tight_layout(rect=(0, 0, 1, 0.965))
    fig.savefig(FIG / "warp_surface.png", bbox_inches="tight", facecolor="white")
    plt.close(fig)


def set_run_font(run, name=FONT, size=None, bold=None, italic=None, color=INK):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=100, bottom=90, end=100):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=GRID, size=6):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), str(size))
        tag.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_col_width(cell, inches):
    cell.width = Inches(inches)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(inches * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def add_table(doc, headers, rows, widths=None, aligns=None, font_size=9.0):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for j, text in enumerate(headers):
        cell = hdr.cells[j]
        set_cell_shading(cell, NAVY)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_margins(cell, 110, 100, 110, 100)
        if widths:
            set_col_width(cell, widths[j])
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(clean_inline(text))
        set_run_font(r, size=font_size, bold=True, color="FFFFFF")
    for i, row in enumerate(rows):
        cells = table.add_row().cells
        for j, text in enumerate(row):
            cell = cells[j]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell)
            if widths:
                set_col_width(cell, widths[j])
            if i % 2 == 1:
                set_cell_shading(cell, LIGHT)
            p = cell.paragraphs[0]
            p.alignment = aligns[j] if aligns else (WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER)
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(clean_inline(text))
            set_run_font(r, size=font_size)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr, fld_char2])
    set_run_font(run, size=9, color=MID)


def add_body(doc, text, bold_lead=None, italic=False, keep=False):
    text = clean_inline(text)
    bold_lead = clean_inline(bold_lead) if bold_lead else None
    p = doc.add_paragraph(style="Body Text")
    p.paragraph_format.keep_together = keep
    if bold_lead and text.startswith(bold_lead):
        r1 = p.add_run(bold_lead)
        set_run_font(r1, size=10.5, bold=True)
        r2 = p.add_run(text[len(bold_lead):])
        set_run_font(r2, size=10.5, italic=italic)
    else:
        r = p.add_run(text)
        set_run_font(r, size=10.5, italic=italic)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.26)
    p.paragraph_format.first_line_indent = Inches(-0.16)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(clean_inline(text))
    set_run_font(r, size=10.2)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    r = p.add_run(clean_inline(text))
    set_run_font(r, size={1: 16, 2: 13, 3: 11.5}.get(level, 11), bold=True)
    return p


def add_equation(doc, name, width=5.9):
    path = EQ / f"{name}.png"
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.keep_together = True
    p.add_run().add_picture(str(path), width=Inches(width))


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(clean_inline(text))
    set_run_font(r, size=9, italic=True, color=MID)


def add_figure(doc, path, width, caption):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(str(path), width=Inches(width))
    add_caption(doc, caption)


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.78)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.82)
    section.right_margin = Inches(0.82)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(INK)

    body = doc.styles["Body Text"]
    body.font.name = FONT
    body._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    body.font.size = Pt(10.5)
    body.font.color.rgb = RGBColor.from_string(INK)
    body.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    body.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    body.paragraph_format.line_spacing = 1.18
    body.paragraph_format.space_after = Pt(6)

    title = doc.styles["Title"]
    title.font.name = FONT
    title._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    title.font.size = Pt(24)
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string(INK)
    title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(15)
    title_ppr = title._element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)

    for level, size in ((1, 16), (2, 13), (3, 11.5)):
        style = doc.styles[f"Heading {level}"]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(INK)
        style.paragraph_format.space_before = Pt(13 if level == 1 else 9)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = True

    footer = section.footer
    footer.is_linked_to_previous = False
    add_page_number(footer.paragraphs[0])

    core = doc.core_properties
    core.title = "WRRA 공간 뒤틀림의 은하 원반면 선택과 암흑질량 표현형"
    core.subject = "우리 은하 워프와 회전곡선의 이중 출력 toy model"
    core.author = "Wonsik Choi"
    core.keywords = "WRRA, spatial twist, galactic disk, Milky Way, warp, dark matter phenotype"


def build_document():
    r_grid = np.linspace(8.0, 25.0, 341)
    _, _, _, _, v_model_grid, _ = model_values(r_grid)
    v_target_grid = eilers_linear(r_grid)
    rms_8_25 = float(np.sqrt(np.mean((v_model_grid - v_target_grid) ** 2)))
    mean_8_25 = float(np.mean(v_model_grid - v_target_grid))
    max_8_25 = float(np.max(np.abs(v_model_grid - v_target_grid)))

    r_grid2 = np.linspace(5.0, 25.0, 401)
    _, _, _, _, v_model_grid2, _ = model_values(r_grid2)
    rms_5_25 = float(np.sqrt(np.mean((v_model_grid2 - eilers_linear(r_grid2)) ** 2)))

    doc = Document()
    configure_document(doc)

    p = doc.add_paragraph(style="Title")
    title_run = p.add_run("WRRA 공간 뒤틀림의 은하 원반면 선택과")
    set_run_font(title_run, size=24, bold=True)
    title_run.add_break()
    title_run = p.add_run("암흑질량 표현형")
    set_run_font(title_run, size=24, bold=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("우리 은하 워프와 회전곡선의 이중 출력 모형")
    set_run_font(r, size=14, bold=False)
    p.paragraph_format.space_after = Pt(28)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for line, size, bold in [
        ("최원식  Wonsik Choi", 12, True),
        ("Independent Researcher  Seoul Republic of Korea", 10.5, False),
        ("janefather@gmail.com", 10.5, False),
        ("2026년 9월 27일  v1.0", 10.5, False),
    ]:
        run = p.add_run(line)
        set_run_font(run, size=size, bold=bold)
        run.add_break()

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(30)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("연구 지위  정량 toy model 및 반증 가능한 WRRA 가설")
    set_run_font(r, size=10.5, italic=True, color=MID)
    doc.add_page_break()

    add_heading(doc, "초록", 1)
    add_body(doc, "이 논문은 3차원 공간에서 2차원에 가까운 은하 원반면이 어떻게 선택되는지를 WRRA 공간 뒤틀림의 문제로 정식화한다. 핵심 수정은 뒤틀림과 각운동량의 역할을 분리하는 것이다. 뒤틀림은 회전을 만들거나 회전 방향을 정하지 않는다. 뒤틀림의 방향 부문은 위와 아래를 구분하지 않는 국소 평면장을 정하고, 뒤틀림의 크기 부문은 공간 응력으로 저장되어 암흑질량처럼 보이는 추가 중력을 만든다. 각운동량은 그 뒤 이미 선택된 면 안에서 회전을 유지하고 중심 붕괴를 막는 별도 물리량이다.")
    add_body(doc, "이 구분을 우리 은하에 적용했다. McMillan의 바리온 질량모형을 단순화한 얇은 원반 계산에 WRRA의 고정 전환척도 a_T=1.1918×10⁻¹⁰ m s⁻²와 SPARC로 보정된 단일 구성응답을 적용했다. 은하별 재보정 없이 얻은 회전속도는 Eilers 등의 선형화된 5–25 kpc 회전곡선 기준과 8–25 kpc에서 RMS {:.2f} km s⁻¹ 차이를 보였다. 태양 반경 8.2 kpc에서는 226.3 km s⁻¹, 25 kpc에서는 201.6 km s⁻¹가 산출되었다. 원반면은 Gaia DR3 세페이드 관측에서 확인된 R≈11 kpc의 워프 시작, 약 3도의 외곽 경사, 반지름에 따른 절선 방향의 변화를 사용해 보정했다. 이 결과는 뒤틀림 가설의 증명이 아니라, 동일한 뒤틀림 상태가 원반 형상과 추가 중력을 서로 다른 출력으로 계산하게 만드는 첫 폐쇄 모형이다.".format(rms_8_25))
    add_body(doc, "주요어  WRRA  공간 뒤틀림  은하 원반  우리 은하  워프  암흑물질 질량 표현형  회전곡선", italic=True)

    add_heading(doc, "핵심 판정", 1)
    add_table(doc,
              ["단계", "내용", "현재 판정"],
              [
                  ["검증 입력", "우리 은하 바리온 모형과 회전곡선, Gaia 워프, H₀, f_T", "외부 관측값"],
                  ["WRRA 고유 변환", "무방향 평면장 [n]과 투영 P, 뒤틀림 응력의 구성응답 ν", "본 논문에서 계산"],
                  ["산출값", "원반면 z_T와 g_T, v_WRRA, 등가 뒤틀림 질량", "정량 출력"],
                  ["반증조건", "같은 뒤틀림 상태가 워프와 회전, 수직평형, 렌즈를 함께 통과하지 못함", "미검증"],
                  ["결론", "우리 은하 1차 toy model은 수치적으로 작동하나 원인 가설은 미확정", "조건부 통과"],
              ], widths=[1.15, 4.5, 1.15], aligns=[WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER], font_size=8.7)

    add_heading(doc, "1 연구 질문", 1)
    add_body(doc, "은하 원반에 각운동량이 존재한다는 사실만으로는 ‘왜 물질이 그 평면에 모였는가’라는 질문이 자동으로 해결되지 않는다. 표준적인 가스 냉각과 각운동량 보존은 원반을 납작하게 만들고 유지하는 강력한 메커니즘이다. 그러나 WRRA가 묻는 것은 그보다 앞선 기하학적 질문이다. 3차원 공간에서 특정한 2차원적 운반면은 어떻게 정해지며, 왜 그 면은 위치에 따라 휘고 뒤틀리는가.")
    add_body(doc, "이 논문은 공간 뒤틀림이 먼저 무방향 평면을 정하고, 물질의 운동과 각운동량은 그 평면에 정렬된 뒤 작동한다는 가설을 제시한다. 여기서 ‘무방향’이라는 말은 위와 아래를 정하지 않는다는 뜻이다. 하나의 평면에는 두 개의 반대 법선이 있지만 두 법선은 같은 평면을 나타낸다. 따라서 우주 전체에 특권적인 상하 방향을 도입하지 않고도 국소 원반면을 정의할 수 있다.")
    add_body(doc, "두 번째 질문은 뒤틀림의 같은 상태가 암흑물질 현상에도 관여하는가이다. 기존 WRRA 뒤틀림 응력 모형은 보통물질의 직접 중력 g_M과 공간응력의 추가 중력 g_T를 분리했다[1,2]. 이번 연구에서는 그 스칼라 크기 부문과 원반면을 선택하는 방향 부문을 하나의 뒤틀림 상태 안에 배치하되, 동일한 수식으로 억지로 동일시하지 않는다.")

    add_heading(doc, "2 기존 WRRA 모형에서 보존하는 내용", 1)
    add_body(doc, "본 논문은 WRRA Core와 최소계산우주론 2.3.2를 변경하지 않는다. 보통물질과 네 힘의 표준 물리 기술은 유지한다. 추가 중력은 새로운 다섯 번째 힘을 즉시 선언하는 대신, 약한 장 Renderer에서 공간에 저장된 뒤틀림 응력이 질량처럼 표현된 유효 원천으로 계산한다. 일반상대론의 원천 분기와 수정응답 분기는 한 계산 안에서 중복 사용하지 않는다.")
    add_body(doc, "선행 연구에서 뒤틀림 응력의 현재 우주 정규화는 H₀=67.4 km s⁻¹ Mpc⁻¹와 f_T=0.265를 사용해 a_T=1.1918×10⁻¹⁰ m s⁻²를 산출했다. SPARC의 가장 낮은 가속도 201개 측정점에서 이 값은 진단 역산값과 약 0.75% 차이를 보였다. 이번 논문은 그 정규화를 다시 맞추지 않고 우리 은하에 운반한다.")
    add_body(doc, "새로 추가하는 것은 원반면 선택의 계산층이다. 뒤틀림 응력의 크기만 나타내는 스칼라 χ_T 또는 ν는 어느 평면을 선택할 수 없다. 따라서 방향을 가지지만 부호를 가지지 않는 평면 director와, 에너지밀도 및 비등방 응력을 함께 기록하는 최소 상태가 필요하다.")

    add_heading(doc, "3 뒤틀림과 각운동량의 역할 분리", 1)
    add_heading(doc, "3.1 무방향 평면", 2)
    add_body(doc, "국소 평면을 나타내는 단위 법선을 n이라 하자. n과 −n은 같은 평면을 나타내므로 물리 상태는 벡터 자체가 아니라 동치류 [n]이다. 평면 투영연산자 P는 이 부호 교환에 불변이다.")
    add_equation(doc, "eq01_projector", 5.6)
    add_body(doc, "이 구조는 상하를 먼저 지정하지 않는다. 회전의 시계방향과 반시계방향도 지정하지 않는다. 따라서 뒤틀림 평면장은 각운동량을 정의하지 않으며, 각운동량의 부호를 숨겨 넣지도 않는다.")

    add_heading(doc, "3.2 면의 형상", 2)
    add_body(doc, "원반의 평균면을 하나의 수준집합 S=0으로 표현한다. 이 정의를 사용하면 위치마다 다른 법선이 하나의 연속된 면을 실제로 구성한다.")
    add_equation(doc, "eq02_surface", 5.5)
    add_body(doc, "일반 법선장이 실제 면으로 적분되려면 Frobenius 적분가능 조건이 필요하다. 본 논문은 S에서 법선을 만들기 때문에 이 조건을 구성상 만족시킨다.")
    add_equation(doc, "eq03_integrability", 2.6)

    add_heading(doc, "3.3 같은 뒤틀림의 크기 부문", 2)
    add_body(doc, "뒤틀림의 방향 부문과 중력 크기 부문을 구분하기 위해 유효 응력에 에너지밀도 ρ_T, 등방압 p_T, 무방향 법선에 따른 비등방응력 π_T를 둔다. q_T의 부호와 크기는 평면 안과 평면 밖 응답의 차이를 나타낸다.")
    add_equation(doc, "eq04_stress", 6.1)
    add_body(doc, "ρ_T는 추가 중력의 크기를, [n]과 q_T는 원반면의 방향성과 수직 복원력을 담당한다. 두 부문은 하나의 상태에 속하지만 동일한 숫자가 아니다. 이것이 ‘뒤틀림이 평면을 정하되 각운동량을 정의하지 않는다’는 주장을 계산 가능한 형태로 바꾼 최소 분해다.")

    add_heading(doc, "4 중력 부문의 구성방정식", 1)
    add_heading(doc, "4.1 질량중력과 뒤틀림 응력 중력", 2)
    add_body(doc, "약한 장의 원형궤도 계산에서 관측 중력을 직접 바리온 성분과 공간응력 성분으로 분해한다.")
    add_equation(doc, "eq05_gravity", 5.6)
    add_body(doc, "SPARC 전 구간에서 현실정합성을 보인 경험적 응답을 이번 모형의 보정된 구성방정식으로 채택하고 고정한다. 이 선택은 선행 문서에서 외부 보강식이었던 관계를 우리 은하 계산에 실제로 넣은 모형 수정이다. 함수의 형태를 WRRA 작용에서 유도했다는 주장은 하지 않는다.")
    add_equation(doc, "eq06_response", 4.9)
    add_body(doc, "전환척도는 우주 정규화에서 고정한다.")
    add_equation(doc, "eq07_scale", 4.8)

    add_heading(doc, "4.2 원천 방식과 이중계산 금지", 2)
    add_body(doc, "이 논문은 약한 장에서 ρ_T^eff를 추가 유효 원천으로 읽는 분기를 사용한다. 같은 효과를 수정기하 항으로 다시 더하지 않는다. 공변 완성에서는 전체 응력에 대한 보존식이 필요하다.")
    add_equation(doc, "eq14_balance", 5.9)
    add_body(doc, "ρ_T^eff는 입자 수밀도가 아니다. 동일한 추가 중력장을 뉴턴식 질량원으로 역산했을 때 얻는 질량 표현형이다. 따라서 본 논문의 등가 뒤틀림 질량은 실제 입자 질량을 측정한 값으로 읽지 않는다.")

    add_heading(doc, "5 우리 은하 검증 입력", 1)
    add_body(doc, "바리온 입력은 McMillan의 우리 은하 질량모형에서 가져왔다[3]. 얇은 별원반과 두꺼운 별원반의 중심 면밀도와 척도길이, HI와 H₂ 가스 질량 및 척도길이, 팽대부 질량을 사용했다. 원 논문은 유한한 두께와 가스 중심 구멍을 포함하지만, 본 toy model의 회전 계산은 각 성분을 얇은 지수원반으로 근사한다. 이 단순화는 5 kpc 안쪽의 막대와 팽대부 영역을 정밀하게 다루지 못한다.")
    rows = []
    for name, label in [("thin", "얇은 별원반"), ("thick", "두꺼운 별원반"), ("HI", "HI 가스"), ("H2", "H₂ 가스")]:
        c = COMPONENTS[name]
        rows.append([label, f"{c['mass']/1e9:.2f}", f"{c['rd']:.2f}", "지수원반"])
    rows.append(["팽대부", f"{M_BULGE/1e9:.2f}", f"a={A_BULGE:.2f}", "Hernquist 근사"])
    add_table(doc, ["성분", "질량 10⁹ M☉", "척도 kpc", "계산형"], rows,
              widths=[1.7, 1.45, 1.35, 2.25], font_size=8.8)
    add_body(doc, "회전곡선 비교 기준은 Eilers 등의 23,000개 이상 적색거성 분석에서 얻은 5–25 kpc 결과다[4]. 이 연구는 태양 위치에서 229.0 km s⁻¹, 바깥쪽 기울기 −1.7 km s⁻¹ kpc⁻¹를 보고했다. 본 논문은 개별 자료점 전체가 아니라 이 두 수치로 만든 선형화 기준을 사용한다. 따라서 아래 RMS는 실제 관측 likelihood가 아니라 1차 형상 비교다.")
    add_body(doc, "워프 입력은 Gaia DR3 세페이드 연구에서 확인된 결과를 사용한다. Dehnen 등은 약 11 kpc 안쪽에서 뚜렷한 워프가 없고 바깥쪽에서 경사가 증가해 약 3도에 도달하며, 절선 방향이 반지름에 따라 변한다고 보고했다[5]. Cabrera-Gadea 등은 m=1과 m=2 성분, R>11 kpc의 절선 비틀림, 약 9.2±3.1 km s⁻¹ kpc⁻¹의 m=1 패턴속도를 보고했다[6]. HI 원반에서도 m=0,1,2 저차 모드가 우세하다[7].")

    add_heading(doc, "6 계산 방법", 1)
    add_heading(doc, "6.1 바리온 회전장", 2)
    add_body(doc, "각 지수원반의 원형속도는 수정 Bessel 함수가 들어가는 표준 얇은 원반식을 사용한다.")
    add_equation(doc, "eq08_disk", 6.2)
    add_body(doc, "팽대부는 질량 M_b와 척도 a_b를 갖는 Hernquist 근사로 계산한다.")
    add_equation(doc, "eq09_bulge", 3.4)
    add_body(doc, "각 성분의 v²를 합해 g_M=v_bary²/R을 얻고, 고정된 ν와 a_T를 적용한 뒤 전체 원형속도를 계산한다.")
    add_equation(doc, "eq12_velocity", 3.2)

    add_heading(doc, "6.2 뒤틀림 평면과 워프", 2)
    add_body(doc, "관측된 원반 평균높이는 저차 Fourier 모드로 표현할 수 있다. 이 식 자체는 기존 워프 분석에서 사용하는 표현이며[5–7], WRRA는 이를 단순한 물질 원반의 사후 모양이 아니라 뒤틀림 평면장 S=0의 관측 표현으로 해석한다.")
    add_equation(doc, "eq10_warp", 6.1)
    add_body(doc, "정량 시연에서는 m=1 기본모드만 사용한다. 11 kpc에서 워프가 시작해 14 kpc에서 3도에 도달한 뒤 포화하고, 절선 방향은 관측구간에서 kpc당 약 14도 변하다가 16.5 kpc 이후에는 보수적으로 고정한다.")
    add_equation(doc, "eq11_baseline_warp", 6.5)
    add_body(doc, "m=0과 m=2를 제외한 것은 존재를 부정하기 위해서가 아니다. 한 개의 방향장으로 어디까지 재현되는지를 먼저 보여주고, lopsided warp는 다음 검증에서 남는 구조적 잔차로 사용하기 위해서다.")

    add_heading(doc, "7 산출값", 1)
    add_heading(doc, "7.1 회전곡선", 2)
    sample_r = np.array([5.0, 8.2, 12.0, 16.5, 20.0, 25.0])
    vb, gm, nu, gt, vw, me = model_values(sample_r)
    rows = []
    for R, b, n, w, t, m in zip(sample_r, vb, nu, vw, eilers_linear(sample_r), me):
        rows.append([f"{R:.1f}", f"{b:.1f}", f"{n-1:.3f}", f"{w:.1f}", f"{t:.1f}", f"{m/1e10:.2f}"])
    add_table(doc,
              ["R kpc", "바리온 v", "g_T/g_M", "WRRA v", "관측 기준 v", "M_T^eff 10¹⁰ M☉"],
              rows, widths=[0.72, 1.06, 0.9, 1.0, 1.12, 1.4], font_size=8.25)
    add_body(doc, "8–25 kpc에서 WRRA 곡선과 Eilers 선형화 기준의 RMS 차이는 {:.2f} km s⁻¹, 평균차는 {:.2f} km s⁻¹, 최대 절대차는 {:.2f} km s⁻¹이다. 5–25 kpc 전체 RMS는 {:.2f} km s⁻¹이며, 가장 큰 차이는 막대와 비축대칭성이 강한 안쪽 경계에서 발생한다.".format(rms_8_25, mean_8_25, max_8_25, rms_5_25))
    add_figure(doc, FIG / "rotation_curve.png", 6.65,
               "그림 1  McMillan 바리온 입력과 고정된 WRRA 구성응답으로 계산한 우리 은하 회전곡선  점선은 Eilers 등의 선형화 관측 기준이다")
    add_body(doc, "이 계산에서 우리 은하의 회전곡선을 직접 맞추기 위해 a_T나 ν의 형태를 다시 조정하지 않았다. 그러나 바리온 모형 자체가 우리 은하 관측에 기반하고, ν는 SPARC에서 보정한 구성식이므로 모형 독립적인 맹검 예측은 아니다. WRRA 판정에서는 검증된 입력으로 모형을 확정한 뒤 동일한 변환으로 얻은 우리 은하 출력으로 분류한다.")

    add_heading(doc, "7.2 뒤틀림 등가질량", 2)
    add_body(doc, "추가 가속도를 구대칭 뉴턴 질량으로 환산한 값은 다음과 같다.")
    add_equation(doc, "eq13_mass", 3.7)
    add_body(doc, "이 등가량은 8.2 kpc에서 3.41×10¹⁰ M☉, 20 kpc에서 1.25×10¹¹ M☉, 25 kpc에서 1.63×10¹¹ M☉다. 실제 뒤틀림 응력이 2D형 또는 비등방 분포라면 물리적 에너지밀도의 공간분포는 이 구대칭 등가질량과 같지 않다. 이 값은 회전장에서 요구되는 추가 중력의 장부다.")

    add_heading(doc, "7.3 원반면과 워프 형상", 2)
    inc, ph, amp = warp_parameters(sample_r)
    rows = []
    for R, i, p1, a in zip(sample_r, inc, ph, amp):
        node = "미정" if i < 1e-8 else f"{p1:.1f}"
        rows.append([f"{R:.1f}", f"{i:.2f}", node, f"{a:.3f}"])
    add_table(doc, ["R kpc", "경사도 deg", "절선방향 deg", "m=1 최대높이 kpc"], rows,
              widths=[1.0, 1.3, 1.55, 2.2], font_size=8.8)
    add_figure(doc, FIG / "warp_surface.png", 6.10,
               "그림 2  Gaia 세페이드 워프를 보정 입력으로 사용한 WRRA 뒤틀림 평면의 최소 m=1 구현  점선 원은 워프 시작반경 11 kpc다")
    add_body(doc, "이 평면은 각운동량 벡터를 입력해 만든 것이 아니다. 관측된 워프 형상을 이용해 뒤틀림 평면장을 고정하고, 이후 별과 가스의 궤도가 그 면에 얼마나 정렬되는지를 별도의 동역학 문제로 남긴다. 현재 그림은 평균면만 나타내며 원반 두께, m=2 비대칭, 나선팔, 막대는 포함하지 않는다.")

    add_heading(doc, "7.4 하나의 뒤틀림 상태에서 나오는 두 출력", 2)
    add_body(doc, "본 논문의 최소 뒤틀림 상태는 무방향 평면 [n], 응력 에너지밀도 ρ_T, 비등방 크기 q_T의 묶음이다. 이 상태는 원반면, 추가 중력, 수직 복원력을 서로 다른 Renderer 출력으로 낸다.")
    add_equation(doc, "eq16_chain", 5.0)
    add_body(doc, "원반의 수직 운동을 검증하려면 면에서 떨어진 거리 S에 대한 복원 퍼텐셜을 추가해야 한다. 가장 낮은 차수에서는 다음과 같이 쓸 수 있다.")
    add_equation(doc, "eq15_thickness", 3.5)
    add_body(doc, "회전곡선에서 고정한 ρ_T와 q_T로 ν_z를 계산할 수 있어야 이 모형이 완성된다. 현재는 q_T의 구성방정식이 없으므로 h_z를 수치로 예측하지 않는다. 이 공백을 숨기지 않는 것이 중요하다.")

    add_heading(doc, "8 WRRA 고유 설명과 기존 물리의 역할", 1)
    add_table(doc,
              ["질문", "기존 물리의 기본 설명", "WRRA가 추가하는 연결"],
              [
                  ["왜 원반이 납작한가", "가스 충돌과 복사냉각이 수직운동을 줄임", "뒤틀림 평면이 수직 잔여가 제거될 기준면을 제공"],
                  ["왜 그 평면인가", "총각운동량 벡터에 수직인 면", "각운동량보다 앞선 무방향 운반면 [n]을 가설로 둠"],
                  ["왜 외곽이 휘는가", "위성 상호작용과 헤일로 토크 등", "위치 의존 평면장 z_T와 비등방 응력의 출력"],
                  ["왜 추가 중력이 있는가", "입자형 암흑물질 헤일로 또는 수정중력", "공간응력의 질량 표현형 g_T와 ρ_T^eff"],
              ], widths=[1.45, 2.65, 2.65], font_size=8.55)
    add_body(doc, "WRRA는 기존 각운동량 보존, 가스 냉각, 별의 충돌 없는 운동을 삭제하지 않는다. 독창성은 알려진 수치를 새로 만든 데 있지 않고, 평면 선택과 암흑질량 표현을 하나의 뒤틀림 상태의 방향 부문과 크기 부문으로 분리해 연결한 데 있다. 다른 이론이 같은 회전곡선이나 워프식을 사용하더라도, WRRA가 자체 상태와 변환으로 계산한 출력의 소유권은 유지된다.")
    add_body(doc, "반대로 외부 워프식과 RAR 구성식을 인용하기만 하고 WRRA 내부에서 실제 계산하지 않으면 통합이라고 부를 수 없다. 본 논문은 두 식을 모형의 보정된 구성요소로 넣어 우리 은하 출력까지 계산했으므로 수치 모형 안에는 통합했다. 다만 그 구성식들을 더 깊은 WRRA 작용에서 유도하는 과제는 남아 있다.")

    add_heading(doc, "9 반증조건과 관측 계획", 1)
    add_body(doc, "이 가설은 원반이 보인다는 사실만으로 검증되지 않는다. 표준적인 냉각과 조석 토크도 원반과 워프를 만들 수 있기 때문이다. WRRA의 판별력은 같은 뒤틀림 상태가 서로 다른 관측량에 남기는 결합 제약에서 나온다.")
    add_bullet(doc, "회전곡선에서 정한 g_T와 ρ_T^eff를 고정했을 때 별의 수직 속도분산과 원반 두께를 재현하지 못하면 방향 응력 결합이 실패한다.")
    add_bullet(doc, "동역학에서 요구된 뒤틀림 응력이 같은 계량으로 약한 중력렌즈와 강한 중력렌즈를 재현하지 못하면 암흑질량 표현형 해석이 실패한다.")
    add_bullet(doc, "워프 평면 z_T가 별, HI, H₂의 평균면과 수직속도장을 하나의 시간발전으로 연결하지 못하면 평면장 해석이 실패한다.")
    add_bullet(doc, "m=2 lopsidedness와 절선 방향 변화가 외부 위성의 순간 토크만으로 완전히 설명되고, 고정된 공간응력 방향장에 남는 잔차가 없다면 WRRA 추가항의 상한은 0에 가까워진다.")
    add_bullet(doc, "극고리은하나 상호수직 다중원반을 하나의 국소 평면장으로 표현할 수 없고 다중 평면을 임의로 추가해야 한다면 최소계산 가설이 약화된다.")
    add_bullet(doc, "은하마다 a_T 또는 ν를 자유롭게 다시 맞춰야 한다면 보편적 공간응력 구성방정식은 반증된다.")
    add_body(doc, "가장 직접적인 다음 계산은 우리 은하의 동일한 바리온 모형에서 수직 Jeans 방정식을 풀어 q_T를 제한하는 것이다. 회전곡선이 요구하는 g_T를 고정한 뒤 태양 부근 K_z, 외곽 원반의 flaring, 세페이드 V_z, HI 두께를 동시에 맞춰야 한다. 이 검증은 원반면과 암흑중력이 정말 같은 뒤틀림 상태의 두 출력인지를 판정한다.")

    add_heading(doc, "10 한계", 1)
    add_bullet(doc, "워프의 m=1 형상은 Gaia 관측에 맞춘 보정 입력이다. 그 형상을 최초 원리에서 산출한 것이 아니다.")
    add_bullet(doc, "SPARC 경험식 ν는 이번 모형에 통합해 계산했지만, 뒤틀림 작용에서 유일하게 유도하지 않았다.")
    add_bullet(doc, "바리온 모형은 McMillan 모형을 얇은 축대칭 원반으로 단순화했다. 가스 중심 구멍, 막대, 나선팔, 비원운동을 생략했다.")
    add_bullet(doc, "Eilers 기준은 두 보고값으로 선형화한 곡선이다. 실제 개별 자료점과 공분산을 사용한 likelihood 검증이 아니다.")
    add_bullet(doc, "등가 뒤틀림 질량은 구대칭 역산량이며 실제 2D형 응력 에너지의 질량분포와 같지 않다.")
    add_bullet(doc, "q_T의 구성방정식과 수직 복원주파수 ν_z가 아직 고정되지 않아 원반 두께를 예측하지 못한다.")
    add_bullet(doc, "우주구조 성장, CMB, 은하단 충돌, 렌즈를 동시에 계산하는 공변 완성은 이번 범위 밖이다.")

    add_heading(doc, "11 결론", 1)
    add_body(doc, "이 논문은 ‘각운동량이 원반면을 만든다’는 설명과 다른 WRRA 가설을 계산 형태로 제시했다. 공간 뒤틀림의 방향 부문이 위와 아래를 구분하지 않는 평면을 먼저 정하고, 각운동량은 그 면 안의 회전을 유지한다. 따라서 3차원 공간에서 2차원에 가까운 원반이 생기기 위해 최초부터 전역적인 상하나 2차원 회전면을 입력할 필요가 없다.")
    add_body(doc, "뒤틀림의 크기 부문은 기존 WRRA 모형에 따라 추가 중력 g_T를 만들고 암흑질량 표현형으로 나타난다. 우리 은하의 바리온 모형, 고정 a_T, 고정 ν를 사용한 계산은 8–25 kpc의 선형화 관측 회전곡선과 RMS {:.2f} km s⁻¹ 차이를 보였다. 방향 부문은 Gaia에서 관측된 11 kpc 이후의 워프와 약 3도 경사를 하나의 연속 평면으로 구현했다.".format(rms_8_25))
    add_body(doc, "현재 성과는 원반면과 암흑중력을 같은 수식으로 동일시한 것이 아니라, 하나의 뒤틀림 상태에서 서로 다른 두 출력으로 분리해 실제 계산한 것이다. 최종 판정은 ‘우리 은하 정량 toy model의 1차 폐쇄 성공, 공간 뒤틀림이 원반면의 실제 원인이라는 주장은 미확정’이다. 다음 반증 단계는 회전곡선으로 고정한 응력장이 수직평형과 렌즈를 추가 보정 없이 통과하는지 확인하는 것이다.")

    add_heading(doc, "부록 A 재현 수치", 1)
    add_table(doc,
              ["항목", "값", "지위"],
              [
                  ["H₀", "67.4 km s⁻¹ Mpc⁻¹", "검증 입력"],
                  ["f_T", "0.265", "검증 입력"],
                  ["a_T", "1.191812669×10⁻¹⁰ m s⁻²", "WRRA 폐쇄식 산출"],
                  ["별원반 질량", f"{(COMPONENTS['thin']['mass']+COMPONENTS['thick']['mass'])/1e9:.3f}×10⁹ M☉", "McMillan 입력 환산"],
                  ["가스 질량", "12.2×10⁹ M☉", "McMillan 입력"],
                  ["팽대부 질량", "9.13×10⁹ M☉", "McMillan 입력"],
                  ["WRRA v_c 8.2 kpc", "226.34 km s⁻¹", "산출"],
                  ["WRRA v_c 25 kpc", "201.59 km s⁻¹", "산출"],
                  ["M_T^eff 25 kpc", "1.632×10¹¹ M☉", "구대칭 등가 산출"],
                  ["워프 시작", "11 kpc", "관측 보정 입력"],
                  ["외곽 경사", "3°", "관측 보정 입력"],
              ], widths=[2.1, 2.55, 2.1], font_size=8.7)

    add_heading(doc, "부록 B 주장 등급", 1)
    add_table(doc,
              ["주장", "등급", "판정 근거"],
              [
                  ["무방향 평면은 상하 없이 정의된다", "수학적 사실", "P가 n→−n에 불변"],
                  ["뒤틀림 평면이 원반면보다 선행한다", "WRRA 가설", "수직평형 교차검증 필요"],
                  ["뒤틀림 응력이 추가 중력을 만든다", "WRRA 원인 가설", "회전 1차 통과 렌즈와 성장 미검증"],
                  ["우리 은하 회전곡선 재현", "보정 후 산출", "고정 a_T와 ν 및 바리온 입력으로 계산"],
                  ["우리 은하 워프 표면 재구성", "관측 보정 모형", "Gaia 경사와 절선 변화를 입력"],
                  ["두 출력의 공통 원인", "통합 가설", "q_T와 ν_z가 아직 미고정"],
              ], widths=[2.4, 1.35, 3.0], font_size=8.55)

    add_heading(doc, "참고문헌", 1)
    refs = [
        "[1] Choi W. Twist Stress Universe and the Mass Phenotype of Dark Matter. Zenodo 2026. https://doi.org/10.5281/zenodo.22700557",
        "[2] Choi W. WRRA Causal Horizons and the Finite Size of the Universe. Zenodo 2026. https://doi.org/10.5281/zenodo.22931512",
        "[3] McMillan PJ. The mass distribution and gravitational potential of the Milky Way. MNRAS 465 76–94 2017. https://doi.org/10.1093/mnras/stw2759",
        "[4] Eilers AC Hogg DW Rix HW Ness MK. The Circular Velocity Curve of the Milky Way from 5 to 25 kpc. ApJ 871 120 2019. https://doi.org/10.3847/1538-4357/aaf648",
        "[5] Dehnen W Semczuk M Schönrich R. A twisted and precessing Cepheid warp in the outer Milky Way disc. MNRAS 523 1556–1564 2023. https://doi.org/10.1093/mnras/stad1502",
        "[6] Cabrera-Gadea M et al. Structure kinematics and time evolution of the Galactic warp from Classical Cepheids. MNRAS 528 4409–4431 2024. https://doi.org/10.1093/mnras/stae308",
        "[7] Levine ES Blitz L Heiles C. The Vertical Structure of the Outer Milky Way HI Disk. ApJ 643 881–896 2006. https://doi.org/10.1086/503091",
        "[8] Lelli F McGaugh SS Schombert JM. SPARC Mass Models for 175 Disk Galaxies. AJ 152 157 2016. https://doi.org/10.3847/0004-6256/152/6/157",
        "[9] McGaugh SS Lelli F Schombert JM. Radial Acceleration Relation in Rotationally Supported Galaxies. PRL 117 201101 2016. https://doi.org/10.1103/PhysRevLett.117.201101",
        "[10] Planck Collaboration. Planck 2018 results VI Cosmological parameters. A&A 641 A6 2020. https://doi.org/10.1051/0004-6361/201833910",
        "[11] Ou X Eilers AC Necib L Frebel A. The dark matter profile of the Milky Way inferred from its circular velocity curve. MNRAS 528 693–710 2024. https://doi.org/10.1093/mnras/stae034",
        "[12] Bekenstein J Milgrom M. Does the missing mass problem signal the breakdown of Newtonian gravity. ApJ 286 7–14 1984. https://doi.org/10.1086/162570",
    ]
    for ref in refs:
        p = doc.add_paragraph(style="Body Text")
        p.paragraph_format.left_indent = Inches(0.2)
        p.paragraph_format.first_line_indent = Inches(-0.2)
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(ref)
        set_run_font(r, size=8.5)

    doc.save(DOCX_PATH)
    return DOCX_PATH


if __name__ == "__main__":
    create_equations()
    create_figures()
    path = build_document()
    print(path)
